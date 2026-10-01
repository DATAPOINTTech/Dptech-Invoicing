"""Extract a best-effort invoice draft from text-based PDF and DOCX files."""
import io
import re
from datetime import datetime
from typing import Any

MAX_IMPORT_SIZE = 10 * 1024 * 1024


def _normalize_text(value: str) -> str:
    if not value:
        return ""
    value = value.replace("\u00a0", " ")
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _number(value: str) -> float | None:
    if value is None:
        return None
    try:
        cleaned = re.sub(r"[^0-9.,-]", "", str(value).replace("(", "-").replace(")", ""))
        if not cleaned or cleaned in {"-", ".", ","}:
            return None
        cleaned = cleaned.replace(",", "")
        return float(cleaned)
    except (TypeError, ValueError):
        return None


def _date(value: str) -> str | None:
    if not value:
        return None
    cleaned = _normalize_text(value).replace(",", "")
    for fmt in ("%d-%m-%y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(cleaned, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _label_value(text: str, labels: tuple[str, ...]) -> str | None:
    labels_re = "|".join(re.escape(label) for label in labels)
    next_label = r"(?:invoice\s*(?:date|no\.?|number)|due\s*date|payment\s*due|(?:bill(?:ed)?\s*to|customer|client|name)|(?:gst|sales\s*tax|tax\s*rate|vat)|(?:payment\s*terms?|terms|notes?|remarks?)|date)"
    match = re.search(
        rf"(?:{labels_re})\s*(?:#|no\.?|number)?\s*[:\-]?\s*"
        rf"(.+?)(?=\s+(?:{next_label})\s*(?:#|no\.?|number)?\s*[:\-]|\n|$)",
        text,
        re.I,
    )
    if not match:
        return None
    value = re.sub(r"^[\s:.-]+|[\s:.-]+$", "", match.group(1).strip())
    return value.rstrip(":") if value else None


def _extract_pdf(content: bytes) -> tuple[str, list[list[str]]]:
    import pdfplumber
    from pypdf import PdfReader

    texts: list[str] = []
    rows: list[list[str]] = []

    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text:
                    texts.append(page_text)
                for table in page.extract_tables() or []:
                    rows.extend([[str(cell or "").strip() for cell in row] for row in table])
    except Exception:
        pass

    if not texts:
        try:
            reader = PdfReader(io.BytesIO(content))
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text:
                    texts.append(page_text)
        except Exception:
            pass

    return _normalize_text("\n".join(texts)), rows


def _extract_docx(content: bytes) -> tuple[str, list[list[str]]]:
    from docx import Document

    document = Document(io.BytesIO(content))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    rows = [[cell.text.strip() for cell in row.cells] for table in document.tables for row in table.rows]
    return _normalize_text("\n".join(paragraphs)), rows


def _row_items(rows: list[list[str]]) -> list[dict[str, Any]]:
    items = []
    for row in rows:
        cells = [cell.strip() for cell in row if cell and cell.strip()]
        if not cells:
            continue
        joined = " ".join(cells).lower()
        if any(word in joined for word in ("description", "subtotal", "grand total", "total amount", "invoice total", "tax", "gst", "discount", "payment", "terms")):
            continue

        numeric_positions = [(idx, _number(cell)) for idx, cell in enumerate(cells) if _number(cell) is not None]
        if len(numeric_positions) < 2:
            continue

        quantity_idx, quantity_value = numeric_positions[0]
        price_idx, price_value = numeric_positions[1]
        if quantity_value is None or price_value is None or quantity_value <= 0 or price_value < 0:
            continue

        quantity = quantity_value
        price = price_value
        description = " ".join(cells[:min(quantity_idx, len(cells))]).strip() or cells[0].strip()
        if len(description) < 2 or re.fullmatch(r"[\W_#0-9.]+", description):
            continue

        if quantity <= 0 or price < 0:
            continue

        item = {"description": description[:500], "quantity": quantity, "unit": "pcs", "unit_price": price}
        if quantity > 100000 or price > 10000000:
            continue
        items.append(item)
    return items


def _text_items(text: str) -> list[dict[str, Any]]:
    items = []
    pattern = re.compile(
        r"^\s*(.+?)\s+"
        r"(\d+(?:\.\d+)?)\s+"
        r"(?:[A-Za-z]+\s+)?"
        r"([$€£₨]?\s*[\d,]+(?:\.\d+)?)"
        r"(?:\s+[$€£₨]?\s*[\d,]+(?:\.\d+)?)?\s*$"
    )
    for line in text.splitlines():
        if re.search(r"\b(?:subtotal|total|tax|gst|discount|amount due|invoice date|due date)\b", line, re.I):
            continue
        match = pattern.match(line)
        if match:
            qty, price = _number(match.group(2)), _number(match.group(3))
            if qty and price is not None:
                items.append({"description": match.group(1).strip()[:500], "quantity": qty, "unit": "pcs", "unit_price": price})
    return items


def parse_invoice_file(filename: str, content: bytes) -> dict[str, Any]:
    if len(content) > MAX_IMPORT_SIZE:
        raise ValueError("File is larger than the 10 MB import limit.")

    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "pdf":
        text, rows = _extract_pdf(content)
    elif extension in {"docx", "doc"}:
        text, rows = _extract_docx(content)
    else:
        raise ValueError("Only PDF and DOCX invoice files are supported.")

    if not text and not rows:
        raise ValueError("No readable invoice data was found. Use a text-based PDF or DOCX file.")

    text = _normalize_text(text)
    date_match = _label_value(text, ("invoice date", "date"))
    due_match = _label_value(text, ("due date", "payment due"))
    tax_match = _label_value(text, ("gst", "sales tax", "tax rate", "vat"))
    client = _label_value(text, ("bill to", "billed to", "customer", "client", "name"))
    if client:
        client = client.split("  ")[0].strip()

    tax_rate = _number(tax_match or "")
    items = _row_items(rows) or _text_items(text)
    return {
        "client_name": client,
        "invoice_date": _date(date_match) if date_match else None,
        "due_date": _date(due_match) if due_match else None,
        "tax_rate": tax_rate if tax_rate is not None and 0 <= tax_rate <= 100 else None,
        "payment_terms": _label_value(text, ("payment terms", "terms")),
        "notes": _label_value(text, ("notes", "remarks")),
        "items": items,
        "warnings": [] if items else ["No line items could be identified. Add them manually before saving."],
    }